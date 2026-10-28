"""Applies the language team's corrections (column E, "Your correction") from the review
workbook back into i18n/<code>.json, then rebuild with: node tools/build-lang.js
Usage: python tools/import-review.py <reviewed.xlsx>"""
import json, os, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEETS = {'Burmese': 'my', 'Hindi': 'hi', 'Tamil': 'ta'}
wb = load_workbook(sys.argv[1], data_only=True)
for sheet, code in SHEETS.items():
    if sheet not in wb.sheetnames: continue
    path = os.path.join(ROOT, 'i18n', code + '.json')
    data = json.load(open(path, encoding='utf-8'))
    changed = 0
    for row in wb[sheet].iter_rows(min_row=2, values_only=True):
        key, fix = row[0], row[4]
        if key and isinstance(fix, str) and fix.strip() and key in data and fix.strip() != data[key]:
            data[key] = fix.strip(); changed += 1
    json.dump(data, open(path, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    print('%s: %d correction(s) applied' % (code, changed))
